"""Create the example template and reference workbooks used by tests.yaml."""

from pathlib import Path

import openpyxl

here = Path(__file__).parent

wb = openpyxl.Workbook()
inp = wb.active
inp.title = "Input"
inp["A1"], inp["A2"], inp["A3"] = "Price", "Quantity", "VAT rate"
inp["B3"] = 0.2
calc = wb.create_sheet("Calc")
calc["A1"], calc["B1"] = "Net", "=Input!B1*Input!B2"
calc["A2"], calc["B2"] = "VAT", "=B1*Input!B3"
calc["A3"], calc["B3"] = "Gross", "=B1+B2"
calc["A4"], calc["B4"] = "Unit price", "=B1/Input!B2"
wb.save(here / "template.xlsx")

ref = openpyxl.Workbook()
ref.active.title = "Expected"
ref["Expected"]["A1"] = 36
ref.save(here / "reference.xlsx")
