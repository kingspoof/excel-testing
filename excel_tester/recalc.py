"""Formula recalculation.

openpyxl cannot evaluate formulas, and when it saves a workbook it drops
all cached formula results. An engine therefore has to compute them:

- "formulas":    pure-Python `formulas` package. No install besides pip,
                 but not every Excel function is supported.
- "libreoffice": headless LibreOffice. Accurate, but must be installed.

Both rewrite the workbook in place so that the computed values can be read
with openpyxl (`data_only=True`).
"""

import re
import shutil
import subprocess
import tempfile
from pathlib import Path


class RecalcError(RuntimeError):
    pass


def find_soffice():
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    for path in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if Path(path).exists():
            return path
    return None


ENGINES = ("formulas", "libreoffice")
_KEY_RE = re.compile(r"^'?\[[^\]]+\](?P<sheet>.+?)'?!(?P<range>[A-Z]+\d+(?::[A-Z]+\d+)?)$")


def recalculate(path, engine="formulas"):
    """Recalculate all formulas in the workbook at `path` (in place)."""
    if engine == "formulas":
        recalculate_formulas(path)
    elif engine == "libreoffice":
        recalculate_libreoffice(path)
    else:
        raise ValueError(f"Unknown engine {engine!r}. Choose one of {ENGINES}.")


def _plain(value):
    """Convert a `formulas` result value to a plain Python value."""
    if hasattr(value, "item"):  # numpy scalar
        value = value.item()
    if type(value).__name__ == "XlError":
        return str(value)  # e.g. "#DIV/0!"
    if isinstance(value, str) and value == "":
        return None
    return value


def recalculate_formulas(path):
    try:
        import formulas
        from openpyxl import load_workbook
        from openpyxl.utils import range_boundaries
    except ImportError as e:
        raise RecalcError("The 'formulas' package is missing. Run: pip install formulas") from e

    path = Path(path)
    try:
        model = formulas.ExcelModel().loads(str(path)).finish()
        solution = model.calculate()
    except Exception as e:  # noqa: BLE001
        raise RecalcError(f"The 'formulas' engine could not calculate the workbook: {e}") from e

    wb = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    sheets = {name.upper(): name for name in wb.sheetnames}
    for key, ranges in solution.items():
        m = _KEY_RE.match(key)
        if not m or m.group("sheet").upper() not in sheets:
            continue  # named ranges, other workbooks, ...
        ws = wb[sheets[m.group("sheet").upper()]]
        min_col, min_row, _, _ = range_boundaries(m.group("range"))
        for r, row in enumerate(ranges.value):
            for c, value in enumerate(row):
                cell = ws.cell(min_row + r, min_col + c)
                # Replace formulas by their results; leave constants alone.
                if cell.data_type == "f" or isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.value = _plain(value)
    wb.save(path)
    wb.close()


def recalculate_libreoffice(path, timeout=120):
    soffice = find_soffice()
    if soffice is None:
        raise RecalcError(
            "LibreOffice was not found. It is needed to calculate formulas. "
            "Install it from https://www.libreoffice.org/ and try again."
        )
    path = Path(path).resolve()
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        # A private profile avoids clashes with a running LibreOffice instance.
        profile = (tmp / "profile").as_uri()
        out_dir = tmp / "out"
        cmd = [
            soffice,
            f"-env:UserInstallation={profile}",
            "--headless",
            "--calc",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out_dir),
            str(path),
        ]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            raise RecalcError(f"LibreOffice did not finish within {timeout}s.")
        result = out_dir / (path.stem + ".xlsx")
        if not result.exists():
            raise RecalcError(
                "LibreOffice could not recalculate the workbook.\n"
                f"stdout: {proc.stdout.strip()}\nstderr: {proc.stderr.strip()}"
            )
        shutil.copyfile(result, path)
