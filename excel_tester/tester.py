"""Run YAML-defined tests against an Excel template."""

import math
import re
import shutil
import tempfile
import traceback
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import yaml

from .recalc import recalculate

EXCEL_ERRORS = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}
_REF_RE = re.compile(r"^(?:'(?P<q>[^']+)'|(?P<s>[^!]+))!\$?(?P<col>[A-Za-z]{1,3})\$?(?P<row>\d+)$")


class SpecError(ValueError):
    """The YAML test specification is invalid."""


@dataclass
class CheckResult:
    description: str
    passed: bool
    message: str = ""


@dataclass
class TestResult:
    __test__ = False  # not a pytest test class
    name: str
    checks: list = field(default_factory=list)
    error: str = ""  # set when the test could not run at all

    @property
    def passed(self):
        return not self.error and all(c.passed for c in self.checks)


def parse_ref(ref):
    """'Sheet 1!B3' or \"'Sheet 1'!B3\" -> ('Sheet 1', 'B3')."""
    m = _REF_RE.match(str(ref).strip())
    if not m:
        raise SpecError(
            f"Invalid cell reference {ref!r}. Use the form 'SheetName!A1', "
            "e.g. 'Input!B3'."
        )
    sheet = m.group("q") or m.group("s")
    return sheet.strip(), f"{m.group('col').upper()}{m.group('row')}"


def get_cell(wb, ref, workbook_label):
    sheet, cell = parse_ref(ref)
    if sheet not in wb.sheetnames:
        raise SpecError(
            f"Sheet {sheet!r} not found in {workbook_label}. "
            f"Available sheets: {', '.join(wb.sheetnames)}"
        )
    return wb[sheet][cell]


def values_equal(actual, expected, tolerance):
    if isinstance(actual, bool) or isinstance(expected, bool):
        return actual == expected
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        return math.isclose(actual, expected, rel_tol=0, abs_tol=tolerance)
    if isinstance(actual, str) and isinstance(expected, str):
        return actual.strip() == expected.strip()
    return actual == expected


class ExcelTester:
    """Tests one Excel template.

    Example:
        tester = ExcelTester("template.xlsx")
        results = tester.run("tests.yaml")
        tester.print_report(results)
    """

    def __init__(self, template, tolerance=1e-9, recalc=True):
        self.template = Path(template).resolve()
        if not self.template.exists():
            raise FileNotFoundError(f"Template not found: {self.template}")
        if self.template.suffix.lower() not in (".xlsx", ".xlsm"):
            raise ValueError(f"Template must be .xlsx or .xlsm, got {self.template.name}")
        self.tolerance = tolerance
        self.recalc = recalc

    # ---------- loading the spec ----------

    def load_spec(self, yaml_path):
        yaml_path = Path(yaml_path).resolve()
        try:
            spec = yaml.safe_load(yaml_path.read_text(encoding="utf-8"))
        except yaml.YAMLError as e:
            raise SpecError(f"{yaml_path.name} is not valid YAML:\n{e}") from e
        if not isinstance(spec, dict) or not isinstance(spec.get("tests"), list):
            raise SpecError(f"{yaml_path.name} must contain a top-level 'tests:' list.")
        return spec, yaml_path.parent

    # ---------- running ----------

    def run(self, yaml_path):
        spec, base_dir = self.load_spec(yaml_path)
        results = []
        for i, test in enumerate(spec["tests"], start=1):
            name = test.get("name", f"Test {i}") if isinstance(test, dict) else f"Test {i}"
            try:
                if not isinstance(test, dict):
                    raise SpecError("Each test must be a mapping with 'inputs' and 'checks'.")
                results.append(self._run_test(name, test, base_dir))
            except SpecError as e:
                results.append(TestResult(name, error=f"Specification error: {e}"))
            except Exception as e:  # noqa: BLE001 - report everything to the user
                results.append(
                    TestResult(
                        name,
                        error=f"{type(e).__name__}: {e}\n{traceback.format_exc(limit=3)}",
                    )
                )
        return results

    def _run_test(self, name, test, base_dir):
        inputs = test.get("inputs") or {}
        checks = test.get("checks") or []
        if not isinstance(inputs, dict):
            raise SpecError("'inputs' must map cell references to values, e.g. 'Input!B1': 10")
        if not isinstance(checks, list) or not checks:
            raise SpecError("'checks' must be a non-empty list.")

        result = TestResult(name)
        with tempfile.TemporaryDirectory() as tmp:
            # 1) copy the template
            copy = Path(tmp) / f"copy{self.template.suffix}"
            shutil.copyfile(self.template, copy)

            # 2) insert data
            keep_vba = copy.suffix.lower() == ".xlsm"
            wb = openpyxl.load_workbook(copy, keep_vba=keep_vba)
            for ref, value in inputs.items():
                cell = get_cell(wb, ref, "the template")
                if isinstance(cell, openpyxl.cell.cell.MergedCell):
                    raise SpecError(f"Input cell {ref} is part of a merged range; use its top-left cell.")
                cell.value = value
            wb.save(copy)
            wb.close()

            if self.recalc:
                recalculate(copy)

            # 3) check values
            result_wb = openpyxl.load_workbook(copy, data_only=True)
            external = {}
            try:
                for check in checks:
                    result.checks.append(self._check(check, result_wb, external, base_dir))
            finally:
                result_wb.close()
                for wb in external.values():
                    wb.close()
        return result

    def _check(self, check, wb, external, base_dir):
        if not isinstance(check, dict) or "cell" not in check:
            raise SpecError(f"Each check needs a 'cell' key, got: {check!r}")
        ref = check["cell"]
        tol = check.get("tolerance", self.tolerance)
        actual = get_cell(wb, ref, "the tested copy").value

        if "equals" in check:
            expected = check["equals"]
            desc = f"{ref} == {expected!r}"
        elif "equals_cell" in check:
            other = check["equals_cell"]
            if "file" in check:
                path = (base_dir / check["file"]).resolve()
                if path not in external:
                    if not path.exists():
                        raise SpecError(f"Reference file not found: {path}")
                    external[path] = openpyxl.load_workbook(path, data_only=True)
                other_cell = get_cell(external[path], other, path.name)
                expected = other_cell.value
                desc = f"{ref} == {path.name} → {other}"
                if expected is None and other_cell.data_type == "f":
                    return CheckResult(desc, False, f"{other} in {path.name} is empty (no saved value).")
            else:
                expected = get_cell(wb, other, "the tested copy").value
                desc = f"{ref} == {other}"
        else:
            raise SpecError(f"Check on {ref} needs either 'equals' or 'equals_cell'.")

        if isinstance(actual, str) and actual in EXCEL_ERRORS:
            return CheckResult(desc, False, f"{ref} shows the Excel error {actual}.")
        if values_equal(actual, expected, tol):
            return CheckResult(desc, True)
        return CheckResult(
            desc,
            False,
            f"expected {expected!r} ({type(expected).__name__}), "
            f"got {actual!r} ({type(actual).__name__})",
        )

    # ---------- reporting ----------

    @staticmethod
    def print_report(results):
        passed = sum(r.passed for r in results)
        for r in results:
            print(f"{'✅' if r.passed else '❌'} {r.name}")
            if r.error:
                for line in r.error.rstrip().splitlines():
                    print(f"     ⚠️  {line}")
                continue
            for c in r.checks:
                mark = "✔" if c.passed else "✘"
                print(f"     {mark} {c.description}" + (f"  →  {c.message}" if c.message else ""))
        print(f"\n{passed}/{len(results)} tests passed.")
